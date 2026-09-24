"""Turns a report context (see app.reports.services.build_report_context)
into downloadable files. Kept separate from routes.py so the PDF/Excel
layout logic doesn't clutter the request-handling code."""
import io
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BRAND_GREEN = colors.HexColor("#0a4d2e")


def _styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoopTitle", fontSize=20, leading=24, textColor=BRAND_GREEN, spaceAfter=4))
    styles.add(ParagraphStyle(name="CoopSubtitle", fontSize=11, textColor=colors.grey, spaceAfter=16))
    styles.add(ParagraphStyle(name="CoopSection", fontSize=13, leading=16, textColor=BRAND_GREEN, spaceBefore=16, spaceAfter=8))
    styles.add(ParagraphStyle(name="CoopBody", fontSize=10, leading=14))
    return styles


def _stat_table(rows, col_widths=(9 * cm, 6 * cm)):
    data = [["Metric", "Value"]] + rows
    t = Table(data, colWidths=list(col_widths))
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BRAND_GREEN),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dddddd")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7f6")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def build_pdf(ctx, report):
    """Return a BytesIO holding the rendered Annual Report PDF."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm,
        title=f"Cooperative Day {ctx['event'].year} Annual Report",
    )
    styles = _styles()
    event = ctx["event"]
    story = []

    story.append(Paragraph("CoopBank — Cooperative Day Annual Report", styles["CoopTitle"]))
    story.append(Paragraph(f"{event.name} · Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", styles["CoopSubtitle"]))

    info_rows = [
        ["Year", str(event.year)],
        ["Theme", event.theme or "—"],
        ["Location", event.location or "—"],
        ["Dates", f"{event.start_date or '—'} to {event.end_date or '—'}"],
        ["Status", event.status],
        ["Coordinator", event.coordinator.full_name if event.coordinator else "—"],
        ["Report Status", report.status if report else "Draft"],
    ]
    story.append(_stat_table(info_rows))

    story.append(Paragraph("Participation", styles["CoopSection"]))
    story.append(_stat_table([
        ["Teams / Committees", ctx["teams_total"]],
        ["Staff Involved", ctx["staff_total"]],
    ]))

    story.append(Paragraph("Tasks", styles["CoopSection"]))
    story.append(_stat_table([
        ["Total Tasks", ctx["tasks_total"]],
        ["Completed", ctx["tasks_completed"]],
        ["Not Completed", ctx["tasks_not_completed"]],
        ["Overdue", ctx["tasks_overdue"]],
        ["Completion Rate", f"{ctx['tasks_completion_rate']}%"],
    ]))

    story.append(Paragraph("Documents &amp; Correspondence", styles["CoopSection"]))
    story.append(_stat_table([
        ["Total Documents", ctx["documents_total"]],
        ["Correspondence Records", ctx["correspondence_total"]],
        ["Correspondence Closed", ctx["correspondence_closed"]],
    ]))

    story.append(Paragraph("Instructions, Decisions &amp; Meetings", styles["CoopSection"]))
    story.append(_stat_table([
        ["Instructions", ctx["instructions_total"]],
        ["Instructions Completed", ctx["instructions_completed"]],
        ["Instructions Overdue", ctx["instructions_overdue"]],
        ["Decisions Recorded", ctx["decisions_total"]],
        ["Meetings Held", ctx["meetings_total"]],
        ["Action Points (Completed / Total)", f"{ctx['action_points_completed']} / {ctx['action_points_total']}"],
    ]))

    story.append(Paragraph("Procurement &amp; Purchases", styles["CoopSection"]))
    story.append(_stat_table([
        ["Purchase Requests", ctx["procurement_total"]],
        ["Completed Purchases", ctx["procurement_completed"]],
        ["Pending Requests", ctx["procurement_pending"]],
        ["Suppliers Used", ctx["suppliers_used"]],
        ["Estimated Cost (Total)", f"{ctx['estimated_cost_total']:,.2f}"],
        ["Paid Amount (Total)", f"{ctx['paid_amount_total']:,.2f}"],
    ]))

    story.append(Paragraph("Activities &amp; Communication", styles["CoopSection"]))
    story.append(_stat_table([
        ["Activities", ctx["activities_total"]],
        ["Activities Completed", ctx["activities_completed"]],
        ["Communication Items", ctx["communications_total"]],
        ["Published", ctx["communications_published"]],
    ]))

    story.append(PageBreak())
    story.append(Paragraph("Narrative Report", styles["CoopTitle"]))

    def narrative_block(title, text):
        story.append(Paragraph(title, styles["CoopSection"]))
        story.append(Paragraph((text or "Not yet documented.").replace("\n", "<br/>"), styles["CoopBody"]))
        story.append(Spacer(1, 6))

    narrative_block("Executive Summary", report.executive_summary if report else None)
    narrative_block("Key Achievements", report.key_achievements if report else None)
    narrative_block("Challenges", report.challenges if report else None)
    narrative_block("Lessons Learned", report.lessons_learned if report else None)
    narrative_block("Recommendations", report.recommendations if report else None)

    doc.build(story)
    buf.seek(0)
    return buf


def _autosize(ws):
    for col in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(length + 2, 12), 45)


def _write_summary_sheet(ws, ctx, report):
    header_fill = PatternFill("solid", fgColor="0A4D2E")
    header_font = Font(color="FFFFFF", bold=True)
    event = ctx["event"]

    ws["A1"] = f"Cooperative Day {event.year} — Annual Report"
    ws["A1"].font = Font(size=14, bold=True, color="0A4D2E")
    ws.merge_cells("A1:B1")
    ws["A2"] = f"Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    ws["A2"].font = Font(italic=True, color="666666")
    ws.merge_cells("A2:B2")

    rows = [
        ("Event Name", event.name), ("Theme", event.theme or "—"),
        ("Location", event.location or "—"),
        ("Start Date", str(event.start_date or "—")), ("End Date", str(event.end_date or "—")),
        ("Event Status", event.status), ("Report Status", report.status if report else "Draft"),
        ("Coordinator", event.coordinator.full_name if event.coordinator else "—"),
        ("", ""),
        ("Teams / Committees", ctx["teams_total"]), ("Staff Involved", ctx["staff_total"]),
        ("Total Tasks", ctx["tasks_total"]), ("Tasks Completed", ctx["tasks_completed"]),
        ("Tasks Not Completed", ctx["tasks_not_completed"]), ("Tasks Overdue", ctx["tasks_overdue"]),
        ("Task Completion Rate (%)", ctx["tasks_completion_rate"]),
        ("Total Documents", ctx["documents_total"]),
        ("Correspondence Records", ctx["correspondence_total"]),
        ("Instructions", ctx["instructions_total"]), ("Instructions Completed", ctx["instructions_completed"]),
        ("Decisions Recorded", ctx["decisions_total"]), ("Meetings Held", ctx["meetings_total"]),
        ("Purchase Requests", ctx["procurement_total"]), ("Completed Purchases", ctx["procurement_completed"]),
        ("Estimated Cost (Total)", ctx["estimated_cost_total"]), ("Paid Amount (Total)", ctx["paid_amount_total"]),
        ("Activities", ctx["activities_total"]), ("Activities Completed", ctx["activities_completed"]),
        ("Communication Items", ctx["communications_total"]), ("Published", ctx["communications_published"]),
    ]
    r = 4
    ws.cell(row=r, column=1, value="Metric").font = header_font
    ws.cell(row=r, column=1).fill = header_fill
    ws.cell(row=r, column=2, value="Value").font = header_font
    ws.cell(row=r, column=2).fill = header_fill
    for label, value in rows:
        r += 1
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=value)

    if report:
        r += 2
        for title, text in [
            ("Executive Summary", report.executive_summary), ("Key Achievements", report.key_achievements),
            ("Challenges", report.challenges), ("Lessons Learned", report.lessons_learned),
            ("Recommendations", report.recommendations),
        ]:
            ws.cell(row=r, column=1, value=title).font = Font(bold=True, color="0A4D2E")
            r += 1
            ws.cell(row=r, column=1, value=text or "Not yet documented.").alignment = Alignment(wrap_text=True)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            r += 2

    _autosize(ws)


def _write_breakdown_sheet(ws, title, rows, headers):
    ws["A1"] = title
    ws["A1"].font = Font(size=12, bold=True, color="0A4D2E")
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0A4D2E")
    for r_idx, row in enumerate(rows, start=4):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    _autosize(ws)


def build_excel(ctx, report):
    """Return a BytesIO holding the rendered Annual Report workbook."""
    wb = Workbook()
    _write_summary_sheet(wb.active, ctx, report)
    wb.active.title = "Summary"

    ws = wb.create_sheet("Tasks by Status")
    _write_breakdown_sheet(ws, "Tasks by Status", [[r["status"], r["count"]] for r in ctx["tasks_by_status"]], ["Status", "Count"])

    ws = wb.create_sheet("Documents by Status")
    _write_breakdown_sheet(ws, "Documents by Status", [[r["status"], r["count"]] for r in ctx["documents_by_status"]], ["Status", "Count"])

    ws = wb.create_sheet("Procurement by Status")
    _write_breakdown_sheet(ws, "Procurement by Status", [[r["status"], r["count"]] for r in ctx["procurement_by_status"]], ["Status", "Count"])

    ws = wb.create_sheet("Activities by Status")
    _write_breakdown_sheet(ws, "Activities by Status", [[r["status"], r["count"]] for r in ctx["activities_by_status"]], ["Status", "Count"])

    ws = wb.create_sheet("Instructions by Status")
    _write_breakdown_sheet(ws, "Instructions by Status", [[r["status"], r["count"]] for r in ctx["instructions_by_status"]], ["Status", "Count"])

    ws = wb.create_sheet("Communication by Type")
    _write_breakdown_sheet(ws, "Communication by Type", [[r["type"], r["count"]] for r in ctx["communications_by_type"]], ["Type", "Count"])

    ws = wb.create_sheet("Staff Participation")
    _write_breakdown_sheet(
        ws, "Staff Participation",
        [[u.full_name, u.department.name if u.department else "—", u.position or "—"] for u in ctx["staff"]],
        ["Full Name", "Department", "Position"],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_comparison_excel(contexts):
    """One workbook comparing several years side by side."""
    from app.reports.services import COMPARISON_METRICS

    wb = Workbook()
    ws = wb.active
    ws.title = "Year Comparison"
    ws["A1"] = "CoopBank Cooperative Day — Year Comparison"
    ws["A1"].font = Font(size=14, bold=True, color="0A4D2E")

    header_row = 3
    ws.cell(row=header_row, column=1, value="Metric").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=header_row, column=1).fill = PatternFill("solid", fgColor="0A4D2E")
    for i, ctx in enumerate(contexts, start=2):
        c = ws.cell(row=header_row, column=i, value=str(ctx["event"].year))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0A4D2E")

    for r_idx, (key, label) in enumerate(COMPARISON_METRICS, start=header_row + 1):
        ws.cell(row=r_idx, column=1, value=label)
        for c_idx, ctx in enumerate(contexts, start=2):
            ws.cell(row=r_idx, column=c_idx, value=ctx[key])

    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
